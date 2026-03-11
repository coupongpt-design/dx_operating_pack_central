import logging
import os
import re
from typing import Dict, Tuple

from openpyxl import load_workbook
from pptx import Presentation
from pptx.util import Pt


logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
)


SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
PREFERRED_XLSX = "교통비 정리.xlsx"
PREFERRED_PPTX = "교통비 정리.pptx"
OUTPUT_PPTX = "결과.pptx"

DATA_START_ROW = 5
NAME_COL = "D"
TOTAL_COL = "K"
RECEIPT_COL = "L"
ACTUAL_COL = "M"

DIRECT_REF_PATTERN = re.compile(r"^\s*=([A-Z]+[0-9]+)\s*$", re.IGNORECASE)
CELL_REF_PATTERN = re.compile(r"\b([A-Z]+[0-9]+)\b")
SUFFIXED_NAME_PATTERN = re.compile(r"^(.*?)([A-Z])$")


def find_input_file(extension: str, preferred_name: str, exclude_names=None) -> str:
    exclude_names = set(exclude_names or [])
    preferred_path = os.path.join(SCRIPT_DIR, preferred_name)
    if os.path.exists(preferred_path):
        return preferred_path

    candidates = []
    for name in os.listdir(SCRIPT_DIR):
        if not name.endswith(extension):
            continue
        if name.startswith("~$") or name in exclude_names:
            continue
        candidates.append(name)

    if not candidates:
        raise FileNotFoundError(f"{extension} 입력 파일을 찾지 못했습니다.")

    candidates.sort()
    return os.path.join(SCRIPT_DIR, candidates[0])


def format_scalar(value) -> str:
    if value in (None, ""):
        return "0"
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def normalize_name(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip()


def split_top_level_args(text: str):
    args = []
    current = []
    depth = 0

    for char in text:
        if char == "," and depth == 0:
            args.append("".join(current).strip())
            current = []
            continue
        if char == "(":
            depth += 1
        elif char == ")":
            depth -= 1
        current.append(char)

    args.append("".join(current).strip())
    return [arg for arg in args if arg]


def expand_sum_functions(expression: str) -> str:
    upper_expression = expression.upper()

    while "SUM(" in upper_expression:
        start = upper_expression.index("SUM(")
        open_index = start + 3
        depth = 0
        close_index = None

        for idx in range(open_index, len(expression)):
            char = expression[idx]
            if char == "(":
                depth += 1
            elif char == ")":
                depth -= 1
                if depth == 0:
                    close_index = idx
                    break

        if close_index is None:
            break

        inner = expression[open_index + 1:close_index]
        args = split_top_level_args(inner)
        replacement = "(" + " + ".join(args) + ")" if args else "0"
        expression = expression[:start] + replacement + expression[close_index + 1:]
        upper_expression = expression.upper()

    return expression


def safe_eval(expression: str):
    allowed = re.compile(r"^[0-9\.\+\-\*\/\(\) ]+$")
    if not allowed.match(expression):
        raise ValueError(f"평가할 수 없는 식입니다: {expression}")
    return eval(expression, {"__builtins__": {}}, {})


def render_expression(formula_sheet, value_sheet, formula_text: str) -> str:
    expression = formula_text.lstrip("=").strip()

    def replace_ref(match):
        coord = match.group(1).upper()
        value = value_sheet[coord].value
        if value in (None, ""):
            value = formula_sheet[coord].value
        return format_scalar(value)

    expression = CELL_REF_PATTERN.sub(replace_ref, expression)
    expression = expand_sum_functions(expression)
    expression = re.sub(r"\s+", " ", expression).strip()
    return expression or "0"


def resolve_source_value(formula_sheet, value_sheet, coord: str, seen=None) -> Tuple[str, object]:
    seen = seen or set()
    coord = coord.upper()

    if coord in seen:
        raise ValueError(f"순환 참조가 있습니다: {coord}")
    seen.add(coord)

    formula_value = formula_sheet[coord].value
    cached_value = value_sheet[coord].value

    if formula_value in (None, ""):
        return "0", 0 if cached_value in (None, "") else cached_value

    if isinstance(formula_value, (int, float)):
        return format_scalar(formula_value), formula_value

    if not isinstance(formula_value, str):
        return str(formula_value), cached_value if cached_value is not None else formula_value

    direct_ref = DIRECT_REF_PATTERN.match(formula_value)
    if direct_ref:
        return resolve_source_value(formula_sheet, value_sheet, direct_ref.group(1), seen)

    expression = render_expression(formula_sheet, value_sheet, formula_value)

    result = cached_value
    if result in (None, ""):
        try:
            result = safe_eval(expression)
        except Exception:
            result = formula_value

    return expression, result


def choose_display_cell(formula_sheet, value_sheet, row_idx: int) -> Tuple[str, object]:
    for col in (ACTUAL_COL, RECEIPT_COL, TOTAL_COL):
        coord = f"{col}{row_idx}"
        formula_value = formula_sheet[coord].value
        cached_value = value_sheet[coord].value

        if formula_value not in (None, "") or cached_value not in (None, ""):
            return resolve_source_value(formula_sheet, value_sheet, coord)

    return "0", 0


def load_excel_data(excel_path: str) -> Dict[str, str]:
    wb_formula = load_workbook(excel_path, data_only=False)
    wb_value = load_workbook(excel_path, data_only=True)

    sheet_formula = wb_formula.active
    sheet_value = wb_value.active

    data = {}

    for row_idx in range(DATA_START_ROW, sheet_formula.max_row + 1):
        name = sheet_formula[f"{NAME_COL}{row_idx}"].value
        if not name:
            continue

        expression, result = choose_display_cell(sheet_formula, sheet_value, row_idx)
        display_text = f"{expression} = {format_scalar(result)}"
        normalized_name = normalize_name(str(name))
        data[normalized_name] = display_text

        hospital = sheet_formula[f"E{row_idx}"].value
        name_match = SUFFIXED_NAME_PATTERN.match(normalized_name)
        if hospital and name_match:
            alias = f"{name_match.group(1)}({str(hospital).strip()})"
            data[normalize_name(alias)] = display_text

    logging.info("엑셀 데이터 %s건 로드", len(data))
    return data


def add_result_box(slide, anchor_shape, text: str):
    gap = 120000
    box = slide.shapes.add_textbox(
        anchor_shape.left,
        anchor_shape.top + anchor_shape.height + gap,
        anchor_shape.width,
        500000,
    )

    frame = box.text_frame
    frame.clear()
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.font.size = Pt(18)


def update_presentation(pptx_path: str, data: Dict[str, str], output_path: str):
    prs = Presentation(pptx_path)
    updated = 0

    for slide in prs.slides:
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue

            name = normalize_name(shape.text)
            if name not in data:
                continue

            add_result_box(slide, shape, data[name])
            updated += 1
            break

    prs.save(output_path)
    logging.info("PPT %s건 업데이트", updated)


def main():
    excel_path = find_input_file(".xlsx", PREFERRED_XLSX)
    pptx_path = find_input_file(".pptx", PREFERRED_PPTX, exclude_names={OUTPUT_PPTX})
    output_path = os.path.join(SCRIPT_DIR, OUTPUT_PPTX)

    logging.info("엑셀 파일: %s", os.path.basename(excel_path))
    logging.info("PPT 파일: %s", os.path.basename(pptx_path))

    data = load_excel_data(excel_path)
    update_presentation(pptx_path, data, output_path)

    print(f"완료: {OUTPUT_PPTX} 생성")


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        logging.exception("실행 실패")
        print(f"Execution failed: {exc}")
