import os
import re
import logging
import sys
from pptx import Presentation
from openpyxl import load_workbook

try:
    script_dir = os.path.dirname(os.path.abspath(sys.executable if getattr(sys, "frozen", False) else __file__))
except:
    script_dir = os.getcwd()

# 콘솔 + 파일 로그 동시 출력
logging.basicConfig(
    level=logging.DEBUG,
    format="%(asctime)s [%(levelname)s] %(message)s",
    handlers=[
        logging.FileHandler(os.path.join(script_dir, "debug_log.txt"), encoding="utf-8"),
        logging.StreamHandler()
    ]
)

def find_file_with_extension(directory, extension):
    preferred_files = {
        ".pptx": "교통비 정리.pptx",
        ".xlsx": "교통비 정리.xlsx",
    }

    preferred_name = preferred_files.get(extension.lower())
    if preferred_name:
        preferred_path = os.path.join(directory, preferred_name)
        if os.path.exists(preferred_path):
            return preferred_path

    for file in os.listdir(directory):
        if (
            file.endswith(extension)
            and not file.startswith("~$")
            and file != "updated_presentation.pptx"
        ):
            return os.path.join(directory, file)
    raise FileNotFoundError(f"No {extension} file found")

def evaluate_formula(sheet, formula):

    if formula is None:
        return 0, 0

    if not isinstance(formula, str):
        return formula, formula

    if formula.startswith("="):
        formula = formula[1:]

    cell_pattern = re.compile(r"([A-Z]+[0-9]+)")

    def replace_cell(match):
        cell_ref = match.group(1)
        cell = sheet[cell_ref]
        value = cell.value

        if isinstance(value, str) and value.startswith("="):
            _, value = evaluate_formula(sheet, value)

        return str(value if value else 0)

    replaced_formula = cell_pattern.sub(replace_cell, formula)

    sum_pattern = re.compile(r"SUM\(([^()]*)\)", re.IGNORECASE)
    while True:
        match = sum_pattern.search(replaced_formula)
        if not match:
            break

        replaced_formula = replaced_formula.replace(
            match.group(0),
            f"({match.group(1)})",
            1
        )

    try:
        result = eval(replaced_formula, {"__builtins__": {}}, {})
    except:
        result = replaced_formula

    return replaced_formula, result


def load_excel_data_with_evaluation(path):

    wb = load_workbook(path, data_only=False)
    sheet = wb.active

    data = {}

    for row in sheet.iter_rows(min_row=5):

        slide_name = row[3].value

        hospital_name = row[4].value
        total_cell = row[10]          # K열
        receipt_cell = row[11]        # L열
        actual_payment_cell = row[12] # M열

        if not slide_name:
            continue

        receipt_value = receipt_cell.value
        actual_value = actual_payment_cell.value

        receipt_col = receipt_cell.column_letter
        total_coord = total_cell.coordinate

        logging.info(f"Slide 처리: {slide_name}")

        formula = None

        # 0️⃣ 실제지급금액이 비어 있으면 제출영수증, 없으면 합계 사용
        if actual_value in (None, ""):

            if receipt_value not in (None, ""):

                logging.info("실제지급금액 없음 → 제출영수증 사용")
                formula = receipt_cell.value

            else:

                logging.info("실제지급금액 없음 → 합계 사용")
                formula = total_cell.value

        # 1️⃣ 실제지급금액이 제출영수증 열을 참조
        elif isinstance(actual_value, str) and re.search(rf"{receipt_col}[0-9]+", actual_value):

            logging.info("실제지급금액이 제출영수증 열 참조 → 제출영수증 사용")
            formula = receipt_cell.value

        # 2️⃣ 값이 동일
        elif receipt_value == actual_value:

            logging.info("값 동일 → 제출영수증 사용")
            formula = receipt_cell.value

        # 3️⃣ 실제지급금액이 합계 참조
        elif isinstance(actual_value, str) and total_coord in actual_value:

            logging.info("합계 참조 → 합계 사용")
            formula = total_cell.value

        # 4️⃣ 기본
        else:

            logging.info("실제지급금액 계산식 사용")
            formula = actual_payment_cell.value

        replaced_formula, result = evaluate_formula(sheet, formula)

        slide_name = str(slide_name).strip()
        result_text = f"{replaced_formula} = {result}"

        data[slide_name] = result_text

        if hospital_name:
            alias_match = re.match(r"^(.*?)([A-Z])$", slide_name)
            if alias_match:
                alias_name = f"{alias_match.group(1)}({str(hospital_name).strip()})"
                data[alias_name] = result_text

    return data


def update_presentation(pptx_path, data, output_path):

    prs = Presentation(pptx_path)

    for slide in prs.slides:

        for shape in slide.shapes:

            if not shape.has_text_frame:
                continue

            text = shape.text.strip()

            if text in data:

                result_height = 180000
                result_top = max(0, shape.top - result_height - 20000)
                tx_box = slide.shapes.add_textbox(
                    left=shape.left,
                    top=result_top,
                    width=shape.width,
                    height=result_height
                )

                tx_box.text_frame.text = data[text]

    prs.save(output_path)


try:

    pptx_path = find_file_with_extension(script_dir, ".pptx")
    excel_path = find_file_with_extension(script_dir, ".xlsx")

    output_path = os.path.join(script_dir, "updated_presentation.pptx")

    excel_data = load_excel_data_with_evaluation(excel_path)

    update_presentation(pptx_path, excel_data, output_path)

    print("완료: updated_presentation.pptx 생성")

except Exception as e:

    logging.error(e)
    print("Execution failed:", e)
