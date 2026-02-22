from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from .csv_loader import load_csv

_EXCEL_EXTENSIONS = {".xlsx", ".xlsm", ".xltx", ".xltm"}
_CSV_EXTENSIONS = {".csv", ".txt", ".tsv"}


def _to_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value)


def _normalize_header_names(raw_headers: List[Any]) -> List[str]:
    headers: List[str] = []
    used: set[str] = set()
    for idx, raw in enumerate(raw_headers, start=1):
        base = _to_text(raw).strip() or f"col{idx}"
        cand = base
        seq = 2
        while cand in used:
            cand = f"{base}_{seq}"
            seq += 1
        used.add(cand)
        headers.append(cand)
    return headers


def _normalize_dict_rows(raw_rows: List[Dict[str, Any]]) -> tuple[List[Dict[str, str]], List[str], int]:
    if not raw_rows:
        return [], [], 0
    raw_headers = list(raw_rows[0].keys())
    headers = _normalize_header_names(raw_headers)

    rows: List[Dict[str, str]] = []
    empty_row_count = 0
    for raw in raw_rows:
        values = [_to_text(raw.get(key)) for key in raw_headers]
        row = {headers[idx]: values[idx] for idx in range(len(headers))}
        if all(not str(v).strip() for v in row.values()):
            empty_row_count += 1
            continue
        rows.append(row)
    return rows, headers, empty_row_count


def _load_excel_rows(file_path: str) -> Tuple[bool, Optional[List[Dict[str, str]]], Optional[str], Dict[str, Any]]:
    try:
        from openpyxl import load_workbook  # type: ignore
    except Exception:
        return False, None, "Excel support requires openpyxl package.", {}

    try:
        workbook = load_workbook(file_path, read_only=True, data_only=True)
    except Exception as exc:
        return False, None, f"Excel read error: {exc}", {}

    try:
        sheet = workbook.active
        iterator = sheet.iter_rows(values_only=True)
        header_row = next(iterator, None)
        if header_row is None:
            info = {"source_type": "excel", "columns": [], "row_count": 0, "empty_row_count": 0}
            return True, [], None, info

        raw_data_rows = list(iterator)
        max_len = len(header_row)
        for row in raw_data_rows:
            if row is not None:
                max_len = max(max_len, len(row))

        header_values = list(header_row) + [None] * max(0, max_len - len(header_row))
        headers = _normalize_header_names(header_values)

        rows: List[Dict[str, str]] = []
        empty_row_count = 0
        for raw in raw_data_rows:
            raw = tuple(raw or ())
            values = [_to_text(raw[idx] if idx < len(raw) else None) for idx in range(max_len)]
            row = {headers[idx]: values[idx] for idx in range(max_len)}
            if all(not str(v).strip() for v in row.values()):
                empty_row_count += 1
                continue
            rows.append(row)

        info = {
            "source_type": "excel",
            "columns": headers,
            "row_count": len(rows),
            "empty_row_count": empty_row_count,
        }
        return True, rows, None, info
    except Exception as exc:
        return False, None, f"Excel parse error: {exc}", {}
    finally:
        try:
            workbook.close()
        except Exception:
            pass


def _load_csv_rows(file_path: str) -> Tuple[bool, Optional[List[Dict[str, str]]], Optional[str], Dict[str, Any]]:
    success, raw_rows, error = load_csv(file_path)
    if not success:
        return False, None, error, {}
    rows, headers, empty_row_count = _normalize_dict_rows(raw_rows or [])
    info = {
        "source_type": "csv",
        "columns": headers,
        "row_count": len(rows),
        "empty_row_count": empty_row_count,
    }
    return True, rows, None, info


def _load_rows_with_info(file_path: str) -> Tuple[bool, Optional[List[Dict[str, str]]], Optional[str], Dict[str, Any]]:
    path = str(file_path or "").strip()
    if not path:
        return False, None, "Data file path is empty.", {}
    ext = Path(path).suffix.lower()
    if ext in _EXCEL_EXTENSIONS:
        return _load_excel_rows(path)
    if ext in _CSV_EXTENSIONS or not ext:
        return _load_csv_rows(path)
    return False, None, f"Unsupported data file extension: {ext}", {}


def load_data_rows(file_path: str) -> Tuple[bool, Optional[List[Dict[str, str]]], Optional[str]]:
    success, rows, error, _ = _load_rows_with_info(file_path)
    return success, rows, error


def inspect_data_file(file_path: str) -> Tuple[bool, Dict[str, Any], Optional[str]]:
    success, rows, error, info = _load_rows_with_info(file_path)
    if not success:
        return False, {"columns": [], "row_count": 0, "empty_row_count": 0}, error
    merged = dict(info)
    merged.setdefault("columns", [])
    merged.setdefault("row_count", len(rows or []))
    merged.setdefault("empty_row_count", 0)
    return True, merged, None
