import logging
import base64
import re
import shutil
import time
import pandas as pd
from pathlib import Path
from typing import List, Dict, Optional
from datetime import datetime
from config import CONFIG

class DataHandler:
    @staticmethod
    def sanitize_filename(name: str, fallback: str) -> str:
        clean = re.sub(r'[\\/:*?"<>|]+', "_", name).strip()
        return clean if clean else fallback

    @staticmethod
    def _extract_extension(src: Optional[str], default_ext: str) -> str:
        if src:
            match = re.search(r"\.([a-zA-Z0-9]{1,5})(?:[?#]|$)", src)
            if match:
                return f".{match.group(1)}"
        return default_ext if default_ext.startswith(".") else f".{default_ext}"

    @staticmethod
    def build_attachment_filename(msg_date: Optional[datetime], msg_index: int, label: str, src: Optional[str], default_ext: str) -> str:
        date_prefix = msg_date.strftime("%Y%m%d_%H%M%S") if msg_date else f"unknown_{msg_index}"
        safe_label = DataHandler.sanitize_filename(label, "file")
        ext = DataHandler._extract_extension(src, default_ext)
        return f"{date_prefix}_{safe_label}{ext}"

    @staticmethod
    def _save_blob_in_chunks(page, url: str, save_path: Path, max_retries: int) -> bool:
        chunk_size = CONFIG.get("BLOB_CHUNK_SIZE", 1024 * 1024)
        try:
            chunk_size = int(chunk_size)
        except (TypeError, ValueError):
            chunk_size = 1024 * 1024
        if chunk_size <= 0:
            chunk_size = 1024 * 1024

        for attempt in range(max_retries):
            cache_info = None
            try:
                cache_info = page.evaluate("""
                    async (targetUrl) => {
                        window.__codex_blob_cache = window.__codex_blob_cache || {};
                        try {
                            const response = await fetch(targetUrl);
                            if (!response.ok) return null;
                            const blob = await response.blob();
                            const id = 'blob_' + Date.now() + '_' + Math.random().toString(16).slice(2);
                            window.__codex_blob_cache[id] = blob;
                            return { id, size: blob.size };
                        } catch (e) { return null; }
                    }
                """, url)

                if not cache_info:
                    raise RuntimeError("blob fetch returned null")

                blob_id = cache_info.get("id")
                blob_size = int(cache_info.get("size", 0))
                if not blob_id or blob_size <= 0:
                    raise RuntimeError("blob cache missing id/size")

                with open(save_path, "wb") as f:
                    offset = 0
                    while offset < blob_size:
                        end = min(offset + chunk_size, blob_size)
                        chunk_b64 = page.evaluate("""
                            ({ id, start, end }) => {
                                const cache = window.__codex_blob_cache;
                                if (!cache || !cache[id]) return null;
                                const blob = cache[id];
                                const slice = blob.slice(start, end);
                                return new Promise((resolve) => {
                                    const reader = new FileReader();
                                    reader.onloadend = () => {
                                        const result = reader.result;
                                        if (typeof result !== 'string') return resolve(null);
                                        const commaIndex = result.indexOf(',');
                                        resolve(commaIndex >= 0 ? result.slice(commaIndex + 1) : null);
                                    };
                                    reader.readAsDataURL(slice);
                                });
                            }
                        """, {"id": blob_id, "start": offset, "end": end})

                        if not chunk_b64:
                            raise RuntimeError("chunk fetch failed")

                        f.write(base64.b64decode(chunk_b64))
                        offset = end

                if save_path.stat().st_size > 0:
                    return True
                logging.warning(f"파일 저장됨 그러나 크기 0: {save_path}")
            except Exception as e:
                logging.warning(f"Chunked blob save failed ({attempt+1}/{max_retries}): {e}")
                try:
                    if save_path.exists():
                        save_path.unlink()
                except Exception:
                    pass
            finally:
                if cache_info and cache_info.get("id"):
                    try:
                        page.evaluate("""
                            (id) => {
                                if (window.__codex_blob_cache) {
                                    delete window.__codex_blob_cache[id];
                                }
                            }
                        """, cache_info.get("id"))
                    except Exception:
                        pass
            time.sleep(1.0)

        return False

    @staticmethod
    def save_blob_to_file(page, url: str, save_path: Path, max_retries: int = 3) -> bool:
        """Blob URL을 파일로 저장 (페이지 컨텍스트 필요)"""
        if not url: return False
        
        # 이미 존재하는 파일이면 스킵 (선택 사항, 여기서는 덮어쓰기 방지)
        if save_path.exists() and save_path.stat().st_size > 0:
            return True

        if url.startswith("data:"):
            try:
                _, encoded = url.split(",", 1)
                save_path.write_bytes(base64.b64decode(encoded))
                if save_path.stat().st_size > 0:
                    return True
            except Exception as e:
                logging.warning(f"Data URL save failed: {e}")
            return False

        if url.startswith("http://") or url.startswith("https://"):
            timeout_ms = CONFIG.get("TIMEOUT_LONG", 60000)
            for attempt in range(max_retries):
                try:
                    # Use direct request for non-blob URLs to reduce in-page memory overhead.
                    response = page.context.request.get(url, timeout=timeout_ms)
                    if response.ok:
                        save_path.write_bytes(response.body())
                        if save_path.stat().st_size > 0:
                            return True
                        logging.warning(f"파일 저장됨 그러나 크기 0: {save_path}")
                    else:
                        logging.warning(f"HTTP download failed ({response.status}): {url}")
                except Exception as e:
                    logging.warning(f"HTTP download failed ({attempt+1}/{max_retries}): {e}")
                    time.sleep(1.0)
            return False

        if url.startswith("blob:"):
            # [Change] Blob URLs are saved in chunks to avoid large base64 allocations.
            return DataHandler._save_blob_in_chunks(page, url, save_path, max_retries)

        for attempt in range(max_retries):
            try:
                # [Reliability] Blob 데이터 추출 전 대기
                page.wait_for_timeout(500)
                
                base64_data = page.evaluate("""
                    async (targetUrl) => {
                        try {
                            const response = await fetch(targetUrl);
                            if (!response.ok) throw new Error('Network response');
                            const blob = await response.blob();
                            return await new Promise((resolve) => {
                                const reader = new FileReader();
                                reader.onloadend = () => resolve(reader.result);
                                reader.readAsDataURL(blob);
                            });
                        } catch (e) { return null; }
                    }
                """, url)
                
                if base64_data:
                    header, encoded = base64_data.split(",", 1)
                    save_path.write_bytes(base64.b64decode(encoded))
                    # [Verification] 파일 크기 검증
                    if save_path.stat().st_size > 0:
                        return True
                    else:
                        logging.warning(f"파일 저장됨 그러나 크기 0: {save_path}")
            except Exception as e:
                logging.warning(f"Blob save failed ({attempt+1}/{max_retries}): {e}")
                time.sleep(1.0) # 재시도 전 대기
        
        return False

    @staticmethod
    def embed_images_to_excel(excel_path: Path, image_folder: Path, chat_data: List[Dict]):
        try:
            from openpyxl import load_workbook
            from openpyxl.drawing.image import Image as XLImage
            from openpyxl.styles import Alignment, PatternFill, Font, Color
            
            wb = load_workbook(str(excel_path))
            ws = wb.active
            
            # [스타일 정의]
            fill_me = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid") # 노란색 (나)
            fill_other = PatternFill(start_color="E7E6E6", end_color="E7E6E6", fill_type="solid") # 회색 (상대)
            font_bold_red = Font(color="FF0000", bold=True)
            font_link = Font(color="0000FF", underline="single")
            
            # 컬럼 너비 설정
            ws.column_dimensions['A'].width = 20
            ws.column_dimensions['B'].width = 15
            ws.column_dimensions['C'].width = 60 # 내용 너비 조정
            ws.column_dimensions['D'].width = 15
            ws.column_dimensions['E'].width = 30
            ws.column_dimensions['F'].width = 15
            ws['E1'] = "이미지"
            ws['F1'] = "원본링크"
            
            keywords = CONFIG.get("HIGHLIGHT_KEYWORDS", [])
            
            for idx, row_data in enumerate(chat_data, start=2):
                # 1. 스타일링 및 키워드 강조
                sender = row_data.get("보낸 사람", "")
                content = row_data.get("내용", "")
                
                # 셀 객체 확보
                cell_sender = ws[f'B{idx}']
                cell_content = ws[f'C{idx}']
                cell_link = ws[f'F{idx}']
                
                # 배경색 및 정렬
                target_fill = fill_me if sender == "나" else fill_other
                align_dir = 'right' if sender == "나" else 'left'
                
                for col in ['A', 'B', 'C', 'D', 'E', 'F']:
                    ws[f'{col}{idx}'].fill = target_fill
                    ws[f'{col}{idx}'].alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
                
                # 내용 컬럼만 정렬 다르게
                cell_content.alignment = Alignment(wrap_text=True, vertical='center', horizontal=align_dir)
                
                # 키워드 강조
                if content:
                    for kw in keywords:
                        if kw in str(content):
                            cell_content.font = font_bold_red
                            break

                # 2. 이미지 및 링크 처리
                saved_files = row_data.get("이미지목록", [])
                if saved_files:
                    target_file = saved_files[0]
                    full_path = image_folder / target_file 
                    
                    if full_path.exists():
                        try:
                            # 링크 추가
                            cell_link.value = "원본 열기"
                            cell_link.hyperlink = str(full_path.resolve())
                            cell_link.font = font_link
                            
                            # 이미지 삽입
                            img = XLImage(str(full_path))
                            if img.height > 200:
                                ratio = 200 / img.height
                                img.height = 200
                                img.width = int(img.width * ratio)
                            
                            img.anchor = f"E{idx}"
                            ws.add_image(img)
                            ws.row_dimensions[idx].height = 165
                        except Exception as e:
                            logging.warning(f"이미지 삽입 실패 ({target_file}): {e}")
            
            # 전체 정렬 루프 제거 (위에서 처리함)
            # for row in ws.iter_rows():
            #     for cell in row:
            #         cell.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
            
            wb.save(str(excel_path))
            logging.info("엑셀 이미지 삽입 및 서식 완료")
        except Exception as e:
            logging.error(f"엑셀 후처리 중 에러: {e}")

    @staticmethod
    def create_template():
        """엑셀 양식 자동 생성"""
        target_path = Path(CONFIG["EXCEL_FILE"])
        if target_path.exists():
            logging.info(f"이미 파일이 존재합니다: {target_path}")
            return

        try:
            df = pd.DataFrame({
                "이름": ["홍길동", "김철수"],
                "전화번호": ["010-1234-5678", "010-9876-5432"],
                "시작일": ["2023-01-01", ""],
                "종료일": ["2023-12-31", ""]
            })
            df.to_excel(target_path, index=False)
            logging.info(f"✅ 엑셀 양식 생성 완료: {target_path.resolve()}")
            print(f"   [알림] '{target_path}' 파일이 생성되었습니다. 내용을 수정 후 다시 실행해주세요.")
        except Exception as e:
            logging.error(f"양식 생성 실패: {e}")

    @staticmethod
    def get_targets() -> Optional[pd.DataFrame]:
        target_path = Path(CONFIG["EXCEL_FILE"])
        if not target_path.exists():
            logging.error(f"파일 없음: {target_path}")
            return None
            
        temp_path = Path("temp_targets_running.xlsx")
        try:
            shutil.copyfile(target_path, temp_path)
            df = pd.read_excel(temp_path, dtype=object)
            
            # [Validation] 필수 컬럼 확인
            required_cols = ["이름", "전화번호"]
            missing = [col for col in required_cols if col not in df.columns]
            if missing:
                logging.error(f"❌ 엑셀 파일에 필수 칸이 없습니다: {', '.join(missing)}")
                print(f"\n[오류] 엑셀 파일({target_path})에 다음 칸이 빠져있습니다: {', '.join(missing)}")
                print("       양식을 다시 만들려면 메뉴에서 '3. 엑셀 양식 생성'을 선택하세요.")
                return None
                
            return df
        except Exception as e:
            logging.error(f"타겟 파일 로딩 실패: {e}")
            return None
        finally:
            if temp_path.exists():
                try: temp_path.unlink()
                except: pass
